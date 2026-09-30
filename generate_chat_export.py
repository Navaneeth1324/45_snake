import os
import json
import re
import html
import pymupdf
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def clean_user_text(text):
    text = re.sub(r'<USER_REQUEST>\s*', '', text)
    text = re.sub(r'\s*</USER_REQUEST>', '', text)
    text = re.sub(r'<ADDITIONAL_METADATA>[\s\S]*?</ADDITIONAL_METADATA>', '', text)
    text = re.sub(r'<USER_SETTINGS_CHANGE>[\s\S]*?</USER_SETTINGS_CHANGE>', '', text)
    text = re.sub(r'\[file "[^"]*" could not be loaded and was omitted\][\s\S]*?==End of PDF==', '', text)
    return text.strip()

def build_exports():
    transcript_path = '/Users/macm2/.gemini/antigravity/brain/cf216ede-16b4-4105-9465-b81455dde678/.system_generated/logs/transcript.jsonl'
    
    dialogue = []
    with open(transcript_path, 'r', encoding='utf-8') as f:
        for line in f:
            try:
                data = json.loads(line)
            except Exception:
                continue
            
            step_type = data.get('type')
            content = data.get('content', '')
            
            if step_type == 'USER_INPUT':
                cleaned = clean_user_text(content)
                if cleaned:
                    dialogue.append(('Student', cleaned))
            elif step_type == 'PLANNER_RESPONSE':
                tool_calls = data.get('tool_calls', [])
                if content and len(content.strip()) > 35 and not tool_calls:
                    # Filter out short transient status messages
                    if not content.startswith("I am checking") and not content.startswith("Still downloading") and not content.startswith("Downloading the"):
                        dialogue.append(('Antigravity AI Assistant', content.strip()))

    print(f"Loaded {len(dialogue)} dialog turns.")

    # 1. Build DOCX
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Lab 4: VibeCoding - Pair Programming Chat History")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(33, 50, 94)

    meta_p = doc.add_paragraph()
    meta_p.add_run("Course Assignment: ").bold = True
    meta_p.add_run("Snake Game (#45 - SETAPESU26/45_snake)\n")
    meta_p.add_run("Tool Used: ").bold = True
    meta_p.add_run("Google Antigravity AI (Pair Programming Assistant)\n")
    meta_p.add_run("Deliverables: ").bold = True
    meta_p.add_run("Code fixes, Before/After 10s gameplay videos, and Chat transcript\n")

    doc.add_heading("Chronological Chat Transcript", level=1)

    for role, text in dialogue:
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

        if role == 'Student':
            set_cell_background(cell, 'EDF2F7')
            badge = p.add_run("👤 STUDENT PROMPT:\n")
            badge.bold = True
            badge.font.color.rgb = RGBColor(43, 108, 176)
            txt_run = p.add_run(text)
            txt_run.font.size = Pt(10.5)
        else:
            set_cell_background(cell, 'F7FAFC')
            badge = p.add_run("🤖 ANTIGRAVITY AI RESPONSE:\n")
            badge.bold = True
            badge.font.color.rgb = RGBColor(40, 167, 69)
            txt_run = p.add_run(text)
            txt_run.font.size = Pt(10.5)

        doc.add_paragraph() # Spacing

    docx_path = '/Users/macm2/.gemini/antigravity/scratch/45_snake/Lab-4/Chat_History_Lab4.docx'
    doc.save(docx_path)
    print(f"Saved DOCX to {docx_path}")

    # 2. Build PDF via PyMuPDF HTML Story
    html_parts = ["""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <style>
    body {
        font-family: Helvetica, Arial, sans-serif;
        color: #2d3748;
        line-height: 1.5;
        font-size: 11pt;
    }
    h1 {
        color: #1a365d;
        font-size: 20pt;
        margin-bottom: 4px;
        border-bottom: 2px solid #3182ce;
        padding-bottom: 6px;
    }
    .meta-box {
        background-color: #ebf8ff;
        border: 1px solid #bee3f8;
        border-radius: 6px;
        padding: 10px 14px;
        margin-bottom: 20px;
        font-size: 10pt;
    }
    .turn {
        border-radius: 6px;
        padding: 12px 14px;
        margin-bottom: 16px;
        page-break-inside: avoid;
    }
    .user-turn {
        background-color: #ebf4ff;
        border-left: 5px solid #3182ce;
    }
    .ai-turn {
        background-color: #f7fafc;
        border-left: 5px solid #38a169;
        border-top: 1px solid #e2e8f0;
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
    }
    .badge {
        font-weight: bold;
        font-size: 10pt;
        text-transform: uppercase;
        margin-bottom: 6px;
        display: block;
    }
    .badge-user { color: #2b6cb0; }
    .badge-ai { color: #2f855a; }
    .content {
        white-space: pre-wrap;
        word-wrap: break-word;
        font-size: 10pt;
    }
    code {
        background: #edf2f7;
        padding: 2px 4px;
        border-radius: 3px;
        font-family: monospace;
    }
    </style>
    </head>
    <body>
    <h1>Lab 4: VibeCoding - Pair Programming Chat Transcript</h1>
    <div class="meta-box">
        <strong>Course:</strong> Computer Science Engineering Lab (Lab 4: VibeCoding)<br>
        <strong>Assignment:</strong> Snake Game (#45 - SETAPESU26/45_snake)<br>
        <strong>Vibe Coding Tool:</strong> Google Antigravity AI Coding Assistant<br>
        <strong>Included Deliverables:</strong> Videos (before/after), Updated Code, and Exported Chat PDF
    </div>
    """]

    for role, text in dialogue:
        escaped = html.escape(text)
        # Convert markdown style code blocks slightly for HTML display
        escaped = re.sub(r'```(.*?)```', r'<code>\1</code>', escaped, flags=re.DOTALL)
        escaped = re.sub(r'`([^`]+)`', r'<code>\1</code>', escaped)

        if role == 'Student':
            html_parts.append(f"""
            <div class="turn user-turn">
                <span class="badge badge-user">👤 Student Prompt</span>
                <div class="content">{escaped}</div>
            </div>
            """)
        else:
            html_parts.append(f"""
            <div class="turn ai-turn">
                <span class="badge badge-ai">🤖 Antigravity AI Assistant</span>
                <div class="content">{escaped}</div>
            </div>
            """)

    html_parts.append("</body></html>")
    full_html = "\n".join(html_parts)

    pdf_path = '/Users/macm2/.gemini/antigravity/scratch/45_snake/Lab-4/Chat_History_Lab4.pdf'
    rect = pymupdf.Rect(40, 40, 555, 780)
    story = pymupdf.Story(html=full_html)
    writer = pymupdf.DocumentWriter(pdf_path)
    more = 1
    while more:
        device = writer.begin_page(pymupdf.Rect(0, 0, 595, 842))
        more, _ = story.place(rect)
        story.draw(device)
        writer.end_page()
    writer.close()
    print(f"Saved PDF to {pdf_path}")

if __name__ == "__main__":
    build_exports()
